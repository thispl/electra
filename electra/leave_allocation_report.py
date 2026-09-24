import frappe
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

from hrms.hr.doctype.leave_application.leave_application import get_leave_details
from hrms.hr.report.employee_leave_balance_summary.employee_leave_balance_summary import execute as run_leave_balance_summary


def _fmt_date(val):
    if not val:
        return ""
    if isinstance(val, str):
        try:
            val = datetime.strptime(val, "%Y-%m-%d").date()
        except Exception:
            return str(val)
    return val.strftime("%d-%m-%Y")


def _safe_sheet_name(name):
    invalid = ['\\', '/', '*', '?', ':', '[', ']']
    for ch in invalid:
        name = name.replace(ch, '')
    return name[:31]


def generate_leave_allocation_report():
    """Generate Excel report with separate worksheet per leave type showing current allocations."""

    today = frappe.utils.today()

    # Get all leave types dynamically
    leave_types = frappe.db.sql_list("SELECT name FROM `tabLeave Type` ORDER BY name")
    if not leave_types:
        print("No leave types found.")
        return

    # Get all active employees
    employees = frappe.db.sql("""
        SELECT
            name,
            employee_number,
            employee_name,
            company,
            branch,
            department,
            designation,
            custom_employee_grade,
            date_of_joining,
            status
        FROM `tabEmployee`
        WHERE status = 'Active'
        ORDER BY employee_name
    """, as_dict=True)

    emp_map = {e["name"]: e for e in employees}

    # Run the Employee Leave Balance Summary report directly to get exact same data
    default_company = frappe.defaults.get_user_default("Company") or frappe.db.get_single_value("Global Defaults", "default_company")

    report_filters = frappe._dict({
        "date": today,
        "company": default_company,
        "employee_status": "Active",
    })

    report_columns, report_data = run_leave_balance_summary(report_filters)

    # Parse report columns to get leave type names (columns after Employee, Employee Name, Department)
    # Column format: "LeaveTypeName:Float:160" or "LeaveTypeName::160"
    report_leave_types = []
    for col in report_columns[3:]:
        lt_name = col.split(":")[0]
        report_leave_types.append(lt_name)

    # Build a map: { employee_id: { leave_type: remaining_leaves } } from report data
    report_balance_map = {}
    for row in report_data:
        emp_id = row[0]
        report_balance_map[emp_id] = {}
        for idx, lt in enumerate(report_leave_types):
            report_balance_map[emp_id][lt] = row[3 + idx]

    # Get current leave allocations for metadata (from_date, to_date, carry_forward)
    allocations = frappe.db.sql("""
        SELECT
            name,
            employee,
            leave_type,
            total_leaves_allocated,
            new_leaves_allocated,
            carry_forwarded_leaves_count,
            from_date,
            to_date
        FROM `tabLeave Allocation`
        WHERE docstatus = 1
          AND expired = 0
          AND %s BETWEEN from_date AND to_date
        ORDER BY leave_type, employee_name
    """, (today,), as_dict=True)

    if not allocations:
        print("No current leave allocations found.")
        return

    # Build allocation metadata map: { (employee, leave_type): alloc }
    alloc_meta_map = {}
    for alloc in allocations:
        alloc_meta_map[(alloc["employee"], alloc["leave_type"])] = alloc

    # Use get_leave_details for full breakdown (total_leaves, leaves_taken) - same as the report
    leave_details_map = {}
    for emp_name in emp_map:
        try:
            details = get_leave_details(emp_name, today)
            leave_details_map[emp_name] = details.get("leave_allocation", {})
        except Exception:
            leave_details_map[emp_name] = {}

    # Group employees by leave type (only those with a non-zero balance or allocation)
    # Use the report's balance as the source of truth for which employees to include
    employees_by_leave_type = {}
    for emp_name in emp_map:
        emp_report = report_balance_map.get(emp_name, {})
        emp_details = leave_details_map.get(emp_name, {})
        for lt in report_leave_types:
            remaining = emp_report.get(lt, 0)
            lt_detail = emp_details.get(lt, {})
            has_alloc = (emp_name, lt) in alloc_meta_map
            # Include employee if they have an allocation or a non-zero balance in the report
            if has_alloc or remaining != 0:
                if lt not in employees_by_leave_type:
                    employees_by_leave_type[lt] = []
                employees_by_leave_type[lt].append((emp_name, remaining, lt_detail))

    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    date_format = "DD-MM-YYYY"

    headers = [
        "Employee ID",
        "Employee Name",
        "Company",
        "Branch",
        "Department",
        "Designation",
        "Grade",
        "Date of Joining",
        "Employee Status",
        "Leave Type",
        "Leave Period (From)",
        "Leave Period (To)",
        "Total Allocated Leaves",
        "Leaves Availed",
        "Current Leave Balance",
        "Carry Forward",
        "Expiry Date",
    ]

    # Create workbook - remove default sheet
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    sheets_created = 0
    for lt in report_leave_types:
        lt_employees = employees_by_leave_type.get(lt, [])
        if not lt_employees:
            continue

        sheet_name = _safe_sheet_name(lt)
        ws = wb.create_sheet(title=sheet_name)
        sheets_created += 1

        # Write headers
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Write data rows
        for row_idx, (emp_name, report_balance, lt_detail) in enumerate(lt_employees, 2):
            emp = emp_map.get(emp_name, {})
            alloc = alloc_meta_map.get((emp_name, lt), {})

            # Use exact values from get_leave_details (same as the report)
            total_allocated = lt_detail.get("total_leaves", alloc.get("total_leaves_allocated", 0) or 0)
            leaves_taken = lt_detail.get("leaves_taken", 0)
            # Use the report's remaining_leaves as the balance (exact match)
            balance = report_balance
            carry_fwd = alloc.get("carry_forwarded_leaves_count", 0) or 0

            row_data = [
                emp_name,
                emp.get("employee_name", ""),
                emp.get("company", ""),
                emp.get("branch", ""),
                emp.get("department", ""),
                emp.get("designation", ""),
                emp.get("custom_employee_grade", ""),
                emp.get("date_of_joining", ""),
                emp.get("status", ""),
                lt,
                alloc.get("from_date", ""),
                alloc.get("to_date", ""),
                total_allocated,
                leaves_taken,
                balance,
                carry_fwd,
                alloc.get("to_date", ""),
            ]

            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border

                # Apply date format to date columns
                if col_idx in (8, 11, 12, 17):
                    if value:
                        cell.value = _fmt_date(value)
                        cell.number_format = date_format

        # Auto-adjust column widths
        for col_idx in range(1, len(headers) + 1):
            max_length = len(str(headers[col_idx - 1]))
            for row_idx in range(2, ws.max_row + 1):
                cell_value = ws.cell(row=row_idx, column=col_idx).value
                if cell_value is not None:
                    max_length = max(max_length, len(str(cell_value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 3, 45)

        # Freeze header row
        ws.freeze_panes = "A2"

        # Apply autofilter to all columns
        last_col = get_column_letter(len(headers))
        last_row = ws.max_row
        ws.auto_filter.ref = f"A1:{last_col}{last_row}"

    # Save
    output_path = "/home/ubuntu/electra-bench/leave_allocation_report.xlsx"
    wb.save(output_path)
    print(f"Report saved to: {output_path}")
    print(f"Total worksheets (leave types with allocations): {sheets_created}")
    print(f"Total active employees: {len(employees)}")
    print(f"Company: {default_company}")
    print(f"Date: {today}")
    for lt in report_leave_types:
        count = len(employees_by_leave_type.get(lt, []))
        if count:
            print(f"  - {lt}: {count} employees")
