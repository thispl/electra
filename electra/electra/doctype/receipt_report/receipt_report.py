# Copyright (c) 2024, Abdulla and contributors
# For license information, please see license.txt

import frappe
from frappe.model.document import Document


from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from frappe.utils import flt, fmt_money, getdate
from six import BytesIO
from dateutil.relativedelta import relativedelta

class ReceiptReport(Document):
	pass

@frappe.whitelist()
def get_receipts_summary():
	filename = 'Receipt Report'
	build_xlsx_response(filename)
	
def make_xlsx(data, sheet_name=None, wb=None, column_widths=None):
	args = frappe.local.form_dict
	start_date = getdate(args.from_date)
	end_date = getdate(args.to_date)
	months = []
	current = start_date
	while current <= end_date:
		months.append(current.strftime("%b-%y"))
		current += relativedelta(months=1)
	
	column_widths = column_widths or []
	if wb is None:
		wb = openpyxl.Workbook()

 
	ws = wb.create_sheet(sheet_name, 0)
	ws.append(["Al - Shaghairi Trading and Contracting Company W.L.L (ELECTRA)"])
	ws.append(["ELECTRA ALL DIVISIONS - RECEIPTS SUMMARY"])
	ws.append([f"Report Period: {start_date} to {end_date}"])
	ws.append([""])
	header = ["Divisions"] + months + ["Total"]

	# Data
	
	ws.append(header)
 
	xlsx_file = BytesIO()
	wb.save(xlsx_file)
	return xlsx_file

def build_xlsx_response(filename):
	xlsx_file = make_xlsx(filename)
	frappe.response['filename'] = filename + '.xlsx'
	frappe.response['filecontent'] = xlsx_file.getvalue()
	frappe.response['type'] = 'binary'


def receipt_report(args):
	data =[]
	total = 0
	ind = 0
	sa = frappe.db.sql("""
		SELECT * 
		FROM `tabPayment Entry`
		WHERE company = %s AND posting_date BETWEEN %s AND %s AND payment_type ='Receive' AND party_type = 'Customer' AND docstatus = 1 order by posting_date  ASC
	""", (args.company, args.from_date, args.to_date), as_dict=True)
	
	journal = frappe.db.sql("""
        SELECT * 
        FROM `tabJournal Entry`
        WHERE 
            company = %s 
            AND posting_date BETWEEN %s AND %s 
            AND docstatus = 1 
            AND voucher_type IN ("Bank Entry", "Cash Entry")
        ORDER BY posting_date ASC
    """, (args.company, args.from_date, args.to_date), as_dict=True)	
	
	for journ in journal:
		doc_journal = frappe.get_all("Journal Entry Account", {"parent": journ.name,"party_type":"customer"}, ["party", "credit_in_account_currency"])
		for c in doc_journal:
			if c.credit_in_account_currency>0:
				ind += 1
				total += c.credit_in_account_currency
				row=[len(data) + 1,journ.posting_date.strftime("%d-%m-%Y"), journ.name, c.party, float(fmt_money(round(c.credit_in_account_currency, 2)).replace(',', '')), '',journ.remarks]
				data.append(row)

	for i in sa:
		document = frappe.get_all("Payment Entry Reference", {"parent": i.name}, ["reference_doctype", "reference_name"])
		sales_person = ''
		if document:
			for j in document:
				if j.reference_doctype == "Sales Order":
					sales_person = frappe.db.get_value("Sales Order", {"name": j.reference_name}, ["sales_person_user"])
				elif j.reference_doctype == "Sales Invoice":
					sales_person = frappe.db.get_value("Sales Invoice", {"name": j.reference_name}, ["sales_person_user"])	
		ind += 1
		total += i.received_amount
		row = [len(data) + 1, i.posting_date.strftime("%d-%m-%Y"), i.name, i.party_name, float(fmt_money(round(i.received_amount, 2)).replace(',', '')), sales_person,i.remarks]
		data.append(row)

	data.append(["Total", "", "", "", float(fmt_money(round(total, 2)).replace(',', '')), '',''])
	return data
	# 	total += i.received_amount
	# row1 =[ind,i.posting_date,i.name,i.party_name,fmt_money(round(i.received_amount, 2)),sales_person]
	# row2 =["Total","","","",fmt_money(round(total, 2)),'']
	# data.append(row1)
	# data.append(row2)
   
