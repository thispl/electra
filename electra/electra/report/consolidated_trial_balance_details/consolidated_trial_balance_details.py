# Copyright (c) 2026, Electra and contributors
# For license information, please see license.txt

import frappe
from frappe import _
from frappe.utils import cint, flt, getdate

GROUP_ORDER = ["ASSET", "Liability", "Equity", "Income Statement"]

ROOT_TYPE_MAP = {
	"Asset": "ASSET",
	"Liability": "Liability",
	"Equity": "Equity",
	"Income": "Income Statement",
	"Expense": "Income Statement",
}

CR_SHEET_NAMES = {
	"Main": "TB-HO",
	"Branch": "TB-Branch",
	"KingFisher": "TB-King Fisher",
	"Marazem": "TB-Marazem",
}

ZERO_CUTOFF = 0.005


def execute(filters=None):
	filters = frappe._dict(filters or {})
	validate_filters(filters)
	columns = get_columns()
	data = build_report_rows(filters)
	return columns, data


def validate_filters(filters):
	if not filters.from_date or not filters.to_date:
		frappe.throw(_("Please set From Date and To Date"))
	if getdate(filters.from_date) > getdate(filters.to_date):
		frappe.throw(_("From Date cannot be after To Date"))

	if filters.fiscal_year:
		fiscal_year = frappe.db.get_value(
			"Fiscal Year", filters.fiscal_year, ["year_start_date", "year_end_date"], as_dict=True
		)
		if fiscal_year:
			filters.year_start_date = getdate(fiscal_year.year_start_date)
			filters.year_end_date = getdate(fiscal_year.year_end_date)

	if not filters.get("year_start_date"):
		filters.year_start_date = getdate(filters.from_date)
	if not filters.get("year_end_date"):
		filters.year_end_date = getdate(filters.to_date)

	# normalize checkbox / multiselect params (strings when passed via GET)
	for flag in (
		"with_period_closing_entry_for_opening",
		"with_period_closing_entry_for_current_period",
		"include_default_book_entries",
		"show_unclosed_fy_pl_balances",
		"show_zero_values",
	):
		filters[flag] = cint(filters.get(flag))

	if isinstance(filters.get("company"), str):
		companies = frappe.parse_json(filters.company)
		filters.company = companies if isinstance(companies, list) else [filters.company]


def get_columns():
	return [
		{"fieldname": "tb_group", "label": _("Groups"), "fieldtype": "Data", "width": 120},
		{"fieldname": "account_head", "label": _("Account Head"), "fieldtype": "Data", "width": 170},
		{"fieldname": "description", "label": _("Description"), "fieldtype": "Data", "width": 340},
		{"fieldname": "division", "label": _("Division"), "fieldtype": "Data", "width": 140},
		{"fieldname": "cr", "label": _("CR"), "fieldtype": "Data", "width": 90},
		{
			"fieldname": "opening",
			"label": _("Opening Balance"),
			"fieldtype": "Float",
			"precision": 2,
			"width": 130,
		},
		{"fieldname": "debit", "label": _("Debit"), "fieldtype": "Float", "precision": 2, "width": 120},
		{"fieldname": "credit", "label": _("Credit"), "fieldtype": "Float", "precision": 2, "width": 120},
		{
			"fieldname": "closing",
			"label": _("Closing Balance"),
			"fieldtype": "Float",
			"precision": 2,
			"width": 130,
		},
	]


def get_company_map():
	"""company -> {division, cr, sort_order}"""
	rows = frappe.get_all(
		"Consolidated TB Company Map",
		fields=["company", "division", "cr_group", "sort_order"],
		order_by="sort_order",
	)
	return {r.company: r for r in rows}


def get_bucket_map():
	"""(account_number, upper(account_name), company_or_empty) -> bucket doc fields"""
	rows = frappe.get_all(
		"Consolidated TB Bucket Account",
		fields=["parent", "account_number", "account_name", "company"],
	)
	buckets = {
		b.name: b
		for b in frappe.get_all(
			"Consolidated TB Bucket",
			fields=["name", "bucket", "account_head", "tb_group", "sort_order"],
			order_by="sort_order",
		)
	}
	mapping = {}
	for r in rows:
		bucket = buckets.get(r.parent)
		if not bucket:
			continue
		key = ((r.account_number or "").strip(), (r.account_name or "").strip().upper(), r.company or "")
		mapping[key] = bucket
	return mapping, buckets


def get_gl_balances(filters, companies):
	"""Aggregated GL per leaf account per company."""
	conditions = ""
	params = {
		"from_date": filters.from_date,
		"to_date": filters.to_date,
		"year_start": filters.year_start_date,
		"companies": tuple(companies),
	}

	if not filters.get("with_period_closing_entry_for_current_period"):
		conditions += " and gle.voucher_type != 'Period Closing Voucher'"

	if filters.get("finance_book"):
		params["finance_book"] = filters.finance_book
		if filters.get("include_default_book_entries"):
			conditions += " and (gle.finance_book = %(finance_book)s or gle.finance_book is null or gle.finance_book = '')"
		else:
			conditions += " and gle.finance_book = %(finance_book)s"

	# Balance sheet accounts: all history before from_date (or is_opening entries).
	# P&L accounts: only entries inside the current fiscal year, unless
	# show_unclosed_fy_pl_balances is requested (standard TB behaviour).
	bs_opening_cond = "(gle.posting_date < %(from_date)s or gle.is_opening = 'Yes')"
	if filters.get("show_unclosed_fy_pl_balances"):
		pl_opening_cond = bs_opening_cond
	else:
		pl_opening_cond = (
			"(gle.posting_date >= %(year_start)s and "
			"(gle.posting_date < %(from_date)s or gle.is_opening = 'Yes'))"
		)

	if not filters.get("with_period_closing_entry_for_opening"):
		pcv_cond = " and gle.voucher_type != 'Period Closing Voucher'"
		bs_opening_cond += pcv_cond
		pl_opening_cond += pcv_cond

	opening_expr = (
		"case when acc.root_type in ('Income', 'Expense') "
		f"then case when {pl_opening_cond} then gle.debit - gle.credit else 0 end "
		f"else case when {bs_opening_cond} then gle.debit - gle.credit else 0 end end"
	)

	return frappe.db.sql(
		"""
		select
			acc.name as account,
			acc.account_number,
			acc.account_name,
			acc.company,
			acc.root_type,
			sum({opening_expr}) as opening,
			sum(case when gle.is_opening = 'No' and gle.posting_date >= %(from_date)s
				then gle.debit else 0 end) as debit,
			sum(case when gle.is_opening = 'No' and gle.posting_date >= %(from_date)s
				then gle.credit else 0 end) as credit
		from `tabAccount` acc
		left join `tabGL Entry` gle
			on gle.account = acc.name and gle.company = acc.company
			and gle.is_cancelled = 0 and gle.posting_date <= %(to_date)s
			{conditions}
		where acc.is_group = 0 and acc.company in %(companies)s
		group by acc.name
		""".format(
			conditions=conditions,
			opening_expr=opening_expr,
		),
		params,
		as_dict=True,
	)


def strip_company_abbr(account_name, company):
	"""Remove the trailing ' - {ABBR}' from an account name."""
	abbr = frappe.get_cached_value("Company", company, "abbr")
	if abbr:
		for sep in (" - ", "-"):
			suffix = sep + abbr
			if account_name.upper().endswith(suffix.upper()):
				return account_name[: -len(suffix)]
	return account_name


def lookup_bucket(mapping, account_number, account_name, company):
	num = (account_number or "").strip()
	name = (account_name or "").strip().upper()
	return mapping.get((num, name, company)) or mapping.get((num, name, ""))


def build_report_rows(filters):
	company_map = get_company_map()
	bucket_map, buckets = get_bucket_map()

	if filters.get("company"):
		companies = list(filters.company)
	else:
		companies = sorted(company_map) if company_map else frappe.get_all("Company", pluck="name")

	gl_rows = get_gl_balances(filters, companies)

	# bucket -> list of leaf rows
	bucket_rows = {}
	unmapped_order = {}
	for d in gl_rows:
		opening = flt(d.opening)
		debit = flt(d.debit)
		credit = flt(d.credit)
		closing = opening + debit - credit

		has_value = any(abs(v) >= ZERO_CUTOFF for v in (opening, debit, credit, closing))
		if not has_value and not filters.get("show_zero_values"):
			continue

		company_info = company_map.get(d.company) or {}
		bucket = lookup_bucket(bucket_map, d.account_number, d.account_name, d.company)

		if bucket:
			bucket_name = bucket.name
		else:
			# unmapped: group under a per-account fallback bucket
			base = strip_company_abbr(
				f"{d.account_number} - {d.account_name}" if d.account_number else d.account_name,
				d.company,
			)
			bucket_name = f"__unmapped__{ROOT_TYPE_MAP.get(d.root_type, '')}__{base.upper()}"
			if bucket_name not in unmapped_order:
				unmapped_order[bucket_name] = {
					"bucket": base,
					"account_head": "Unmapped",
					"tb_group": ROOT_TYPE_MAP.get(d.root_type, "ASSET"),
					"sort_order": 999999,
				}

		row = {
			"description": d.account,
			"division": company_info.get("division") or d.company,
			"cr": company_info.get("cr_group") or "",
			"company": d.company,
			"company_sort": company_info.get("sort_order") or 9999,
			"opening": opening,
			"debit": debit,
			"credit": credit,
			"closing": closing,
			"row_type": "account",
		}
		bucket_rows.setdefault(bucket_name, []).append(row)

	data = []
	for group in GROUP_ORDER:
		# buckets belonging to this group, in mapping order then unmapped
		ordered = []
		for b in bucket_rows:
			if b.startswith("__unmapped__"):
				if unmapped_order[b]["tb_group"] == group:
					ordered.append((unmapped_order[b]["sort_order"], b, unmapped_order[b]))
			elif buckets.get(b) and buckets[b].tb_group == group:
				ordered.append((buckets[b].sort_order or 0, b, buckets[b]))
		ordered.sort(key=lambda x: x[0])

		for _sort, bucket_name, meta in ordered:
			rows = bucket_rows[bucket_name]
			rows.sort(key=lambda r: (r["company_sort"], r["description"]))

			sub = {"opening": 0.0, "debit": 0.0, "credit": 0.0, "closing": 0.0}
			for r in rows:
				for f in sub:
					sub[f] += r[f]

			data.append(
				{
					"tb_group": meta["tb_group"],
					"account_head": meta["account_head"],
					"description": meta["bucket"],
					"division": "All",
					"cr": "",
					"opening": sub["opening"],
					"debit": sub["debit"],
					"credit": sub["credit"],
					"closing": sub["closing"],
					"row_type": "bucket",
					"bold": 1,
				}
			)
			data.extend(rows)

	total = {"opening": 0.0, "debit": 0.0, "credit": 0.0, "closing": 0.0}
	for r in data:
		if r["row_type"] == "account":
			for f in total:
				total[f] += r[f]

	data.append({})
	data.append(
		{
			"tb_group": "",
			"account_head": "",
			"description": "",
			"division": "TOTAL",
			"cr": "",
			"opening": total["opening"],
			"debit": total["debit"],
			"credit": total["credit"],
			"closing": total["closing"],
			"row_type": "total",
			"bold": 1,
		}
	)

	return data


@frappe.whitelist()
def export_excel(**kwargs):
	filters = frappe._dict(kwargs)
	validate_filters(filters)
	data = build_report_rows(filters)

	from openpyxl import Workbook
	from openpyxl.styles import Font
	from openpyxl.utils import get_column_letter

	wb = Workbook()
	bold = Font(bold=True)

	title = "CONSOLIDATED TRIAL BALANCE"
	dated = f"DATED: {getdate(filters.from_date).strftime('%d/%m/%Y')} TO {getdate(filters.to_date).strftime('%d/%m/%Y')}"
	headers = ["Groups ", "Account Head", "Description", "Division", "CR", "Opening Balance", "Debit", "Credit", "Closing Balance"]

	def write_sheet(ws, rows, sl_no=False):
		ncols = len(headers) + (1 if sl_no else 0)
		ws.merge_cells(start_row=2, start_column=1 if not sl_no else 2, end_row=2, end_column=ncols)
		ws.cell(row=2, column=1 if not sl_no else 2, value=title).font = bold
		ws.merge_cells(start_row=3, start_column=1 if not sl_no else 2, end_row=3, end_column=ncols)
		ws.cell(row=3, column=1 if not sl_no else 2, value=dated).font = bold
		hdr = (["Sl No:"] if sl_no else []) + headers
		for i, h in enumerate(hdr, start=1):
			ws.cell(row=5, column=i, value=h).font = bold
		rno = 6
		sl = 1
		for r in rows:
			if not r:
				rno += 1
				continue
			col = 1
			if sl_no:
				ws.cell(row=rno, column=1, value=sl)
				sl += 1
				col = 2
			values = [r.get("tb_group"), r.get("account_head"), r.get("description"), r.get("division"), r.get("cr"),
				r.get("opening"), r.get("debit"), r.get("credit"), r.get("closing")]
			for i, v in enumerate(values):
				c = ws.cell(row=rno, column=col + i, value=v)
				if r.get("bold"):
					c.font = bold
				if i >= 5 and isinstance(v, (int, float)):
					c.number_format = "#,##0.00"
			rno += 1

	ws = wb.active
	ws.title = "Consolidated TB-Details"
	write_sheet(ws, data)

	for cr, sheet_name in CR_SHEET_NAMES.items():
		sub = [r for r in data if r and (r.get("cr") == cr or r.get("row_type") in ("bucket", "total"))]
		# keep bucket subtotal rows only if they have leaf rows in this CR group
		filtered = []
		for i, r in enumerate(sub):
			if r.get("row_type") == "bucket":
				j = i + 1
				keep = False
				while j < len(sub) and sub[j] and sub[j].get("row_type") == "account":
					keep = True
					j += 1
				if keep:
					# recompute subtotal for this CR group only
					s = {"opening": 0.0, "debit": 0.0, "credit": 0.0, "closing": 0.0}
					for k in range(i + 1, j):
						for f in s:
							s[f] += sub[k][f]
					r = dict(r)
					r.update(s)
					filtered.append(r)
			else:
				filtered.append(r)
		# drop grand total row and recompute for the CR group
		filtered = [r for r in filtered if r.get("row_type") != "total"]
		t = {"opening": 0.0, "debit": 0.0, "credit": 0.0, "closing": 0.0}
		for r in filtered:
			if r.get("row_type") == "account":
				for f in t:
					t[f] += r[f]
		filtered.append({})
		filtered.append({"division": "TOTAL", "row_type": "total", "bold": 1, **t})
		write_sheet(wb.create_sheet(sheet_name), filtered, sl_no=True)

	import io

	output = io.BytesIO()
	wb.save(output)
	frappe.response["type"] = "binary"
	frappe.response["filecontent"] = output.getvalue()
	frappe.response["filename"] = "Consolidated Trial Balance.xlsx"
	frappe.response["content_type"] = (
		"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
	)
